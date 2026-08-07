import pandas as pd
from docx import Document
from docx.shared import Pt
import os
import re

# Specific exam rolls to generate
TARGET_ROLLS = [
    54225100,
    52125138,
    71125136,
    11125037,
    71125137,
    71125138,
    54225101,
    74225139,
    64125151,
    44125079
]

def format_value(key: str, value) -> str:
    """Format values based on the field type."""
    if pd.isna(value):
        return ""
    
    value_str = str(value)
    
    # Remove .0 from roll numbers and class
    if key.lower() in ['exam roll', 'roll', 'class']:
        if value_str.endswith('.0'):
            value_str = value_str[:-2]
    
    # Add +88 prefix to mobile numbers if not present
    if 'mobile' in key.lower():
        value_str = value_str.strip().replace(' ', '').replace('+', '')
        if value_str.isdigit():
            if not value_str.startswith('88'):
                value_str = '+88' + value_str
            else:
                value_str = '+' + value_str
    
    return value_str


def replace_placeholders_in_doc(doc: Document, data_row: pd.Series) -> None:
    """Replace all placeholders in a docx Document with values from data_row."""
    placeholder_map = {}
    for key, value in data_row.items():
        placeholder_lower = f"{{{{{key}}}}}".lower()
        formatted_value = format_value(key, value)
        placeholder_map[placeholder_lower] = formatted_value
    
    print(f"\nProcessing Exam Roll: {data_row.get('Exam Roll', 'N/A')}")
    print(f"  Name: {data_row.get('Name (English)', 'N/A')}")

    # Replace in paragraphs
    for p in doc.paragraphs:
        for run in p.runs:
            original_text = run.text
            for placeholder_lower, value in placeholder_map.items():
                pattern = re.compile(re.escape(placeholder_lower), re.IGNORECASE)
                if pattern.search(original_text):
                    run.text = pattern.sub(value, run.text)
                    run.font.name = 'Times New Roman'
                    run.font.bold = False
                    run.font.size = Pt(12)

    # Replace in tables
    for table in doc.tables:
        for row_cells in table.rows:
            for cell in row_cells.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        original_text = run.text
                        for placeholder_lower, value in placeholder_map.items():
                            pattern = re.compile(re.escape(placeholder_lower), re.IGNORECASE)
                            if pattern.search(original_text):
                                run.text = pattern.sub(value, run.text)
                                run.font.name = 'Times New Roman'
                                run.font.bold = False
                                run.font.size = Pt(12)


def main():
    template = r"c:\Users\Lenovo\Downloads\Telegram Desktop\ASP Admit Card Soft Copy (5).docx"
    data_file = r"c:\Users\Lenovo\Downloads\Telegram Desktop\orders-2025-10-14-17-12-00.xlsx"
    output_dir = "Generated_Admit_Cards_Specific_10"
    
    if not os.path.isfile(template):
        print(f"Template file not found: {template}")
        return
    
    if not os.path.isfile(data_file):
        print(f"Data file not found: {data_file}")
        return
    
    os.makedirs(output_dir, exist_ok=True)
    
    # Read Excel file
    print(f"Reading Excel file...")
    df = pd.read_excel(data_file)
    
    print(f"Total rows in Excel: {len(df)}")
    print(f"Target rolls to process: {len(TARGET_ROLLS)}")
    print(f"Target rolls: {TARGET_ROLLS}\n")
    
    # Filter for specific exam rolls
    df_filtered = df[df['Exam Roll'].isin(TARGET_ROLLS)]
    
    print(f"Filtered rows found: {len(df_filtered)}")
    
    if len(df_filtered) == 0:
        print("ERROR: No matching exam rolls found in the Excel file!")
        print(f"Available rolls: {sorted(df['Exam Roll'].tolist())}")
        return
    
    print("\n" + "="*70)
    print("GENERATING ADMIT CARDS FOR SPECIFIC ROLLS")
    print("="*70)
    
    for idx, row in df_filtered.iterrows():
        doc = Document(template)
        replace_placeholders_in_doc(doc, row)
        
        roll = str(row.get("Exam Roll", idx + 1))
        safe_roll = "".join(c for c in roll if c.isalnum() or c in ("-", "_")) or str(idx + 1)
        output_docx = os.path.join(output_dir, f"Admit_{safe_roll}.docx")
        doc.save(output_docx)
        print(f"  ✅ Saved: {output_docx}")
    
    print(f"\n{'='*70}")
    print(f"✅ Generated {len(df_filtered)} admit cards")
    print(f"📁 Location: {os.path.abspath(output_dir)}")
    print(f"{'='*70}")
    
    # List missing rolls
    found_rolls = set(df_filtered['Exam Roll'].tolist())
    missing_rolls = [r for r in TARGET_ROLLS if r not in found_rolls]
    
    if missing_rolls:
        print(f"\n⚠️  WARNING: {len(missing_rolls)} roll(s) not found in Excel:")
        for roll in missing_rolls:
            print(f"  - {roll}")


if __name__ == "__main__":
    main()
