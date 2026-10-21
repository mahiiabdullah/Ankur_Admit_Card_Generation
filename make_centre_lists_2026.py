"""
Build Centre_Wise_Roll_List_2026.xlsx from Admit_Card_Data_2026.xlsx.

  Summary   - students per centre / gender / class
  Centre 1-4 - one sheet per exam centre: a Boys table, then a Girls table,
               each sorted by exam roll, with a serial number. Print-ready.

Students on the 'Separate' sheet (no card) are not included.
"""
import argparse

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

CENTRES = [  # (code, sheet name, centre name)
    ("1", "Centre 1 - Collegiate", "চট্টগ্রাম কলেজিয়েট স্কুল"),
    ("2", "Centre 2 - Patenga", "পতেঙ্গা উচ্চ বিদ্যালয়"),
    ("3", "Centre 3 - Silver Bells", "সিলভার বেলস কিন্ডারগার্টেন এন্ড গার্লস হাই স্কুল"),
    ("4", "Centre 4 - Monsurabad", "নতুন মনছুরাবাদ আলহাজ্ব মোস্তফা হাকিম হাই স্কুল"),
]
COLUMNS = [  # (header, data column, width)
    ("SL", None, 5),
    ("Exam Roll", "Exam Roll", 11),
    ("Student Name", "Name (English)", 30),
    ("Father's Name", "Fathers Name", 28),
    ("Class", "Class", 6),
    ("School", "Name of School", 42),
    ("Mobile", "Mobile (Online)", 16),
    ("Order ID", "Order ID", 9),
]
CLASS_ORDER = ["4", "5", "6", "7", "8", "9", "10"]

THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="1F6F43")
BOYS_FILL = PatternFill("solid", fgColor="DDEBF7")
GIRLS_FILL = PatternFill("solid", fgColor="FCE4EC")
ZEBRA = PatternFill("solid", fgColor="F7F7F7")


def write_table(ws, row, title, df, band_fill):
    ncol = len(COLUMNS)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
    c = ws.cell(row, 1, f"{title}  ({len(df)})")
    c.font = Font(bold=True, size=12)
    c.fill = band_fill
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 20
    row += 1
    for j, (head, _, _) in enumerate(COLUMNS, 1):
        h = ws.cell(row, j, head)
        h.font = Font(bold=True, color="FFFFFF")
        h.fill = HEAD_FILL
        h.alignment = Alignment(horizontal="center", vertical="center")
        h.border = BORDER
    head_row = row
    row += 1
    for i, (_, r) in enumerate(df.iterrows(), 1):
        for j, (_, col, _) in enumerate(COLUMNS, 1):
            v = i if col is None else r[col]
            if col in ("Exam Roll", "Class", "Order ID"):
                v = int(v)
            cell = ws.cell(row, j, v)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=col == "Name of School",
                                       horizontal="center" if col in (None, "Exam Roll", "Class", "Order ID") else "left")
            if i % 2 == 0:
                cell.fill = ZEBRA
        row += 1
    return row + 1, head_row


def main(src, out):
    df = pd.read_excel(src, sheet_name="Admit_Data", dtype=str)
    df = df.dropna(subset=["Exam Roll"])
    df["_roll"] = df["Exam Roll"].astype(int)

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"

    rows = []
    for code, sheet, centre in CENTRES:
        cdf = df[df["Exam Roll"].str[1] == code].sort_values("_roll")
        mismatched = cdf[cdf["Exam Centre"] != centre]
        if len(mismatched):
            raise SystemExit(f"Roll centre digit and Exam Centre disagree for orders {mismatched['Order ID'].tolist()}")
        ws = wb.create_sheet(sheet)
        ncol = len(COLUMNS)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
        t = ws.cell(1, 1, f"Aunkur Scholarship Project 2026 — Centre {code}: {centre}")
        t.font = Font(bold=True, size=14)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
        boys, girls = cdf[cdf["Gender"] == "Boy"], cdf[cdf["Gender"] == "Girl"]
        ws.cell(2, 1, f"Exam: 23 October 2026, 10:00 AM   |   Total {len(cdf)}  (Boys {len(boys)}, Girls {len(girls)})").font = Font(italic=True)
        r = 4
        r, boys_head = write_table(ws, r, "BOYS", boys, BOYS_FILL)
        r, _ = write_table(ws, r, "GIRLS", girls, GIRLS_FILL)
        for j, (_, _, w) in enumerate(COLUMNS, 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = ws.cell(boys_head + 1, 1)
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "1:2"
        ws.oddFooter.center.text = "Page &P of &N"

        for gender, g in (("Boy", boys), ("Girl", girls)):
            counts = g["Class"].value_counts()
            rows.append([f"Centre {code}", centre, gender] + [int(counts.get(c, 0)) for c in CLASS_ORDER] + [len(g)])

    head = ["Centre", "Centre name", "Gender"] + [f"Class {c}" for c in CLASS_ORDER] + ["Total"]
    summary.append(["Aunkur Scholarship Project 2026 — students per centre"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([])
    summary.append(head)
    for j in range(1, len(head) + 1):
        c = summary.cell(3, j)
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF"), HEAD_FILL, BORDER
        c.alignment = Alignment(horizontal="center")
    for rr in rows:
        summary.append(rr)
    total = ["All", "", ""] + [sum(r[3 + i] for r in rows) for i in range(len(CLASS_ORDER) + 1)]
    summary.append(total)
    for row in summary.iter_rows(min_row=4, max_row=summary.max_row):
        for c in row:
            c.border = BORDER
            if row[2].value == "Girl":
                c.fill = GIRLS_FILL
            elif row[2].value == "Boy":
                c.fill = BOYS_FILL
    for c in summary[summary.max_row]:
        c.font = Font(bold=True)
    for j, w in enumerate([10, 46, 8] + [9] * len(CLASS_ORDER) + [8], 1):
        summary.column_dimensions[get_column_letter(j)].width = w

    wb.save(out)
    print(f"Saved {out}: {len(df)} students")
    for rr in rows:
        print(f"  {rr[0]} {rr[2]:<4} {rr[-1]}")
    print(f"  total {total[-1]}")


if __name__ == "__main__":
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--data", default="Admit_Card_Data_2026.xlsx")
    p.add_argument("--out", default="Centre_Wise_Roll_List_2026.xlsx")
    a = p.parse_args()
    main(a.data, a.out)
