from docx import Document
import pandas as pd

# Check a file from the new batch
doc_path = r"Generated_Admit_Cards_Fixed\Admit_41125059.docx"
doc = Document(doc_path)

print("="*70)
print("VERIFYING MOBILE NUMBER IN NEWLY GENERATED FILE")
print("="*70)
print(f"File: {doc_path}\n")

# Extract all text and look for mobile number
all_text = []
for para in doc.paragraphs:
    if para.text.strip():
        all_text.append(para.text)

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                if para.text.strip():
                    all_text.append(para.text)

# Print lines with relevant data
print("Looking for key fields:\n")
for line in all_text:
    if any(keyword in line for keyword in ['Name', 'Father', 'School', 'Class', 'Mobile', 'Exam Roll', 'Centre', '01']):
        print(line)

# Compare with Excel data
print("\n" + "="*70)
print("EXCEL DATA FOR EXAM ROLL 41125059:")
print("="*70)
df = pd.read_excel(r'c:\Users\Lenovo\Downloads\Telegram Desktop\ASP Online Reg. Boy.xlsx', header=1)
student = df[df['Exam Roll'] == 41125059].iloc[0]

print(f"Name (English): {student['Name (English)']}")
print(f"Fathers Name: {student['Fathers Name']}")
print(f"Name of School: {student['Name of School']}")
print(f"Class: {student['Class']}")
print(f"Mobile (Online): {student['Mobile (Online)']}")
print(f"Exam Roll: {student['Exam Roll']}")
print(f"Exam Centre: {student['Exam Centre']}")
