"""
Export all mobile numbers to a single XML file
"""
import os
from docx import Document
from pathlib import Path
import re
import xml.etree.ElementTree as ET
from xml.dom import minidom

def extract_mobile_from_docx(docx_path):
    """Extract mobile number from a DOCX file"""
    try:
        doc = Document(docx_path)
        mobile_pattern = r'(\+?88)?[-\s]?0?1[3-9]\d{8,9}'
        
        all_text = []
        for para in doc.paragraphs:
            all_text.append(para.text)
        
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        all_text.append(para.text)
        
        full_text = ' '.join(all_text)
        
        mobile_numbers = []
        for match in re.finditer(mobile_pattern, full_text):
            number = match.group(0).strip()
            mobile_numbers.append(number)
        
        if mobile_numbers:
            return mobile_numbers[0]
        
        return None
        
    except Exception as e:
        return None


def create_mobile_xml(directory, output_file):
    """Create XML file with all mobile numbers"""
    
    dir_path = Path(directory)
    
    if not dir_path.exists():
        print(f"Directory not found: {directory}")
        return
    
    # Get all DOCX files
    docx_files = list(dir_path.glob("*.docx"))
    
    if not docx_files:
        print(f"No DOCX files found in {directory}")
        return
    
    print(f"Found {len(docx_files)} DOCX files")
    print(f"Extracting mobile numbers...\n")
    
    # Create root element
    root = ET.Element("admit_cards")
    root.set("total_files", str(len(docx_files)))
    
    found_count = 0
    not_found_count = 0
    
    for docx_file in sorted(docx_files):
        # Extract exam roll from filename
        filename = docx_file.stem
        exam_roll = filename.replace('Admit_', '')
        
        mobile = extract_mobile_from_docx(docx_file)
        
        # Create student element
        student = ET.SubElement(root, "student")
        student.set("exam_roll", exam_roll)
        
        # Add filename
        file_elem = ET.SubElement(student, "filename")
        file_elem.text = docx_file.name
        
        # Add mobile number
        mobile_elem = ET.SubElement(student, "mobile_number")
        if mobile:
            mobile_elem.text = mobile
            mobile_elem.set("status", "found")
            found_count += 1
            print(f"✅ {exam_roll}: {mobile}")
        else:
            mobile_elem.text = "NOT_FOUND"
            mobile_elem.set("status", "missing")
            not_found_count += 1
            print(f"❌ {exam_roll}: NOT FOUND")
    
    # Add summary as attributes to root
    root.set("mobile_found", str(found_count))
    root.set("mobile_missing", str(not_found_count))
    
    # Convert to pretty XML string
    xml_string = ET.tostring(root, encoding='unicode')
    dom = minidom.parseString(xml_string)
    pretty_xml = dom.toprettyxml(indent="  ")
    
    # Remove extra blank lines
    pretty_xml = '\n'.join([line for line in pretty_xml.split('\n') if line.strip()])
    
    # Save to file
    with open(output_file, 'w', encoding='utf-8') as f:
        f.write(pretty_xml)
    
    # Print summary
    print(f"\n{'='*70}")
    print(f"XML FILE CREATED")
    print(f"{'='*70}")
    print(f"Total files processed: {len(docx_files)}")
    print(f"Mobile numbers found: {found_count}")
    print(f"Mobile numbers NOT found: {not_found_count}")
    print(f"\n📁 XML file saved to:")
    print(f"  {os.path.abspath(output_file)}")
    print(f"\n📊 File size: {os.path.getsize(output_file):,} bytes")
    print(f"{'='*70}")


if __name__ == "__main__":
    directory = r"C:\Users\Public\Excel to Admit Card\Generated_Admit_Cards_Fixed"
    output_file = os.path.join(directory, "all_mobile_numbers.xml")
    
    create_mobile_xml(directory, output_file)
    
    print(f"\n✅ Done! You can open the XML file to view all mobile numbers.")
