import pandas as pd
from docx import Document

# Check the first DOCX file
doc_path = r"Generated_Admit_Cards\Admit_11125018.docx"
doc = Document(doc_path)

print("=" * 70)
print("SAMPLE ADMIT CARD CONTENT (First generated file)")
print("=" * 70)
print(f"File: {doc_path}\n")

# Extract and print all text
for para in doc.paragraphs:
    if para.text.strip():
        print(para.text)

# Check tables too
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                if para.text.strip():
                    print(para.text)

print("\n" + "=" * 70)
print("Checking Excel data for comparison")
print("=" * 70)

# Read Excel and show first row
df = pd.read_excel(r'c:\Users\Lenovo\Downloads\Telegram Desktop\ASP Online Reg. Boy.xlsx', header=1)
first_row = df.iloc[0]

print(f"\nExam Roll: {first_row['Exam Roll']}")
print(f"Name (English): {first_row['Name (English)']}")
print(f"Fathers Name: {first_row['Fathers Name']}")
print(f"Name of School: {first_row['Name of School']}")
print(f"Class: {first_row['Class']}")
print(f"Mobile (Online): {first_row['Mobile (Online)']}")
print(f"Exam Centre: {first_row['Exam Centre']}")
