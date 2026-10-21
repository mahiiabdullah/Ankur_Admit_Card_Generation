"""
Assign 2026 exam rolls to students in Admit_Card_Data_2026.xlsx that do not
have one yet, following ROLL_NUMBER_RULES.md:

    roll = class code + centre code + gender code + "26" + 3-digit serial

Each (class, centre, gender) group continues from its latest issued roll in
section 4 of ROLL_NUMBER_RULES.md. Students in a group get consecutive serials
in Order ID order. Afterwards the section 4 table and its date are updated.

Rows that already have an Exam Roll are never changed, so re-running is safe.
"""
import argparse
import re
import sys
from datetime import date

from openpyxl import load_workbook

CLASS_CODE = {4: "4", 5: "5", 6: "6", 7: "7", 8: "8", 9: "9", 10: "1"}
CENTRE_CODE = {
    "চট্টগ্রাম কলেজিয়েট স্কুল": "1",
    "পতেঙ্গা উচ্চ বিদ্যালয়": "2",
    "সিলভার বেলস কিন্ডারগার্টেন এন্ড গার্লস হাই স্কুল": "3",
    "নতুন মনছুরাবাদ আলহাজ্ব মোস্তফা হাকিম হাই স্কুল": "4",
}
CENTRES = "1234"
GENDER_CODE = {"boy": "1", "girl": "2"}
YEAR = "26"
UNKNOWN = "—"  # no roll known yet for a group (its starting number must come from the user)

_CELL = rf"\s*(\d{{8}}|{UNKNOWN})\s*\|"
TABLE_ROW_RE = re.compile(r"^\|\s*(\d+)\s*\|\s*(Boy|Girl)\s*\|" + _CELL * len(CENTRES) + r"\s*$")


def read_latest(rules_text: str) -> dict:
    """Return {(class, centre_code, gender_code): latest_roll or None} from section 4."""
    latest = {}
    for line in rules_text.splitlines():
        m = TABLE_ROW_RE.match(line)
        if m:
            cls, gender = int(m.group(1)), GENDER_CODE[m.group(2).lower()]
            for centre, roll in zip(CENTRES, m.groups()[2:]):
                latest[(cls, centre, gender)] = None if roll == UNKNOWN else roll
    expected = len(CLASS_CODE) * len(CENTRES) * 2
    if len(latest) != expected:
        sys.exit(f"Expected {expected} counters in section 4 of the rules file, found {len(latest)}")
    return latest


def write_latest(rules_text: str, latest: dict) -> str:
    out = []
    for line in rules_text.splitlines():
        m = TABLE_ROW_RE.match(line)
        if m:
            cls, gender_word = int(m.group(1)), m.group(2)
            g = GENDER_CODE[gender_word.lower()]
            rolls = " | ".join(latest[(cls, c, g)] or f"{UNKNOWN:<8}" for c in CENTRES)
            line = f"| {cls:<2} | {gender_word:<4} | {rolls} |"
        elif line.startswith("Last updated:"):
            line = f"Last updated: {date.today():%d/%m/%Y}"
        out.append(line)
    return "\n".join(out) + "\n"


def main(xlsx: str, rules: str, dry_run: bool) -> None:
    with open(rules, encoding="utf-8") as f:
        rules_text = f.read()
    latest = read_latest(rules_text)

    wb = load_workbook(xlsx)
    ws = wb["Admit_Data"]
    col = {c.value: c.column - 1 for c in ws[1]}
    rows = list(ws.iter_rows(min_row=2))

    used = {str(r[col["Exam Roll"]].value) for r in rows if r[col["Exam Roll"]].value}
    todo, problems = [], []
    for r in rows:
        if r[col["Exam Roll"]].value:
            continue
        oid = r[col["Order ID"]].value
        cls = r[col["Class"]].value
        centre = CENTRE_CODE.get(r[col["Exam Centre"]].value or "")
        gender = GENDER_CODE.get(str(r[col["Gender"]].value or "").strip().lower())
        try:
            cls = int(cls)
        except (TypeError, ValueError):
            cls = None
        if cls not in CLASS_CODE or not centre or not gender:
            problems.append(f"Order {oid}: class={r[col['Class']].value!r} "
                            f"centre={r[col['Exam Centre']].value!r} gender={r[col['Gender']].value!r}")
            continue
        todo.append((int(oid), cls, centre, gender, r))

    if problems:
        print("Cannot assign rolls - class, centre or gender is missing/unclear for:")
        print("\n".join("  " + p for p in problems))
        sys.exit(1)

    assigned = []
    for oid, cls, centre, gender, r in sorted(todo, key=lambda t: t[0]):
        key = (cls, centre, gender)
        if latest[key] is None:
            sys.exit(f"No starting roll known for class {cls}, centre {centre}, gender {gender} "
                     f"(order {oid}) - ask the user and fill section 4 of the rules file.")
        serial = int(latest[key][-3:]) + 1
        if serial > 999:
            sys.exit(f"Serial would pass 999 for class {cls}, centre {centre}, gender {gender} - stopping.")
        roll = f"{CLASS_CODE[cls]}{centre}{gender}{YEAR}{serial:03d}"
        assert len(roll) == 8 and roll not in used, roll
        used.add(roll)
        latest[key] = roll
        assigned.append((oid, cls, centre, gender, roll, r))

    print(f"Rolls to assign: {len(assigned)}")
    by_group = {}
    for oid, cls, centre, gender, roll, _ in assigned:
        by_group.setdefault((cls, centre, gender), []).append(roll)
    for (cls, centre, gender), rolls in sorted(by_group.items()):
        g = "Boy" if gender == "1" else "Girl"
        print(f"  Class {cls:>2}  Centre {centre}  {g:<4}  {len(rolls):>3}  {rolls[0]} - {rolls[-1]}")

    if dry_run:
        print("\nDry run - nothing written.")
        return

    for oid, cls, centre, gender, roll, r in assigned:
        r[col["Exam Roll"]].value = roll
        fname = r[col["File Name"]].value or ""
        r[col["File Name"]].value = f"Admit_{roll}" + ("_Duplicate" if fname.endswith("_Duplicate") else "")
    wb.save(xlsx)

    with open(rules, "w", encoding="utf-8") as f:
        f.write(write_latest(rules_text, latest))

    print(f"\nSaved rolls to {xlsx} and updated section 4 of {rules}")


if __name__ == "__main__":
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--xlsx", default="Admit_Card_Data_2026.xlsx")
    p.add_argument("--rules", default="ROLL_NUMBER_RULES.md")
    p.add_argument("--dry-run", action="store_true", help="Show what would be assigned without saving")
    a = p.parse_args()
    main(a.xlsx, a.rules, a.dry_run)
