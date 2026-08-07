import argparse
import sys
import pandas as pd
from docx import Document
try:
    from docx2pdf import convert
    _HAS_DOCX2PDF = True
except Exception:
    _HAS_DOCX2PDF = False
import os
import re
from docx.shared import Pt


def format_value(key: str, value) -> str:
    """Format values based on the field type."""
    if pd.isna(value):
        return ""
    
    value_str = str(value)
    
    # Remove .0 from roll numbers and class
    if key.lower() in ['exam roll', 'roll', 'class']:
        # Remove .0 if it's a float that ends with .0
        if value_str.endswith('.0'):
            value_str = value_str[:-2]
    
    # Add +88 prefix to mobile numbers if not present
    if 'mobile' in key.lower():
        # Remove any existing spaces, + symbols
        value_str = value_str.strip().replace(' ', '').replace('+', '')
        # Only add +88 if it's a number and doesn't start with 88
        if value_str.isdigit():
            if not value_str.startswith('88'):
                value_str = '+88' + value_str
            else:
                value_str = '+' + value_str
    
    return value_str


def replace_placeholders_in_doc(doc: Document, data_row: pd.Series) -> None:
    """Replace all placeholders in a docx Document with values from data_row.

    Placeholders are expected in the form {{ColumnName}}.
    Performs case-insensitive matching and formats text as Times New Roman, not bold.
    """
    # Create a mapping of lowercase placeholder to actual value
    placeholder_map = {}
    for key, value in data_row.items():
        # Handle both {{key}} formats - normalize the placeholder
        placeholder_lower = f"{{{{{key}}}}}".lower()
        formatted_value = format_value(key, value)
        placeholder_map[placeholder_lower] = formatted_value
    
    print(f"\nProcessing row with Exam Roll: {data_row.get('Exam Roll', 'N/A')}")
    print(f"Data available: {dict(data_row)}")

    # Replace in paragraphs
    for p in doc.paragraphs:
        for run in p.runs:
            original_text = run.text
            for placeholder_lower, value in placeholder_map.items():
                # Case-insensitive replacement
                pattern = re.compile(re.escape(placeholder_lower), re.IGNORECASE)
                if pattern.search(original_text):
                    print(f"  Replacing {placeholder_lower} with {value}")
                    run.text = pattern.sub(value, run.text)
                    # Format: Times New Roman, not bold
                    run.font.name = 'Times New Roman'
                    run.font.bold = False
                    run.font.size = Pt(12)

    # Replace in tables
    for table in doc.tables:
        for row_cells in table.rows:
            for cell in row_cells.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        original_text = run.text
                        for placeholder_lower, value in placeholder_map.items():
                            # Case-insensitive replacement
                            pattern = re.compile(re.escape(placeholder_lower), re.IGNORECASE)
                            if pattern.search(original_text):
                                print(f"  Replacing {placeholder_lower} with {value}")
                                run.text = pattern.sub(value, run.text)
                                # Format: Times New Roman, not bold
                                run.font.name = 'Times New Roman'
                                run.font.bold = False
                                run.font.size = Pt(12)


def main(template: str, data_file: str, output_dir: str, header_row: int = 0):
    # Normalize and check paths
    template = os.path.expanduser(template)
    data_file = os.path.expanduser(data_file)
    output_dir = os.path.expanduser(output_dir)

    if not os.path.isfile(template):
        print(f"Template file not found: {template}")
        sys.exit(2)
    if not os.path.isfile(data_file):
        print(f"Data file not found: {data_file}")
        sys.exit(2)

    os.makedirs(output_dir, exist_ok=True)

    # Read Excel file
    print(f"Reading Excel file with header at row {header_row}...")
    df = pd.read_excel(data_file, header=header_row)
    
    print(f"Columns found: {list(df.columns)}")
    print(f"Total rows: {len(df)}\n")

    for idx, row in df.iterrows():
        doc = Document(template)
        replace_placeholders_in_doc(doc, row)

        roll = str(row.get("Exam Roll", idx + 1))
        safe_roll = "".join(c for c in roll if c.isalnum() or c in ("-", "_")) or str(idx + 1)
        output_docx = os.path.join(output_dir, f"Admit_{safe_roll}.docx")
        doc.save(output_docx)
        print(f"Saved: {output_docx}")

    print(f"\nGenerated {len(df)} admit cards")
    
    if _HAS_DOCX2PDF:
        print("\nConverting to PDF...")
        try:
            convert(output_dir)
            print("PDF conversion complete!")
        except Exception as e:
            print(f"Warning: failed to convert docx to pdf: {e}")
    else:
        print("docx2pdf not available; skipping conversion to PDF.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Generate admit cards from an Excel file and a Word template.")
    parser.add_argument("--template", required=True, help="Path to the .docx template")
    parser.add_argument("--data", required=True, help="Path to the Excel file with data")
    parser.add_argument("--out", default="Generated_Admit_Cards", help="Output directory")
    parser.add_argument("--header", type=int, default=0, help="Row number for header (0-indexed)")
    args = parser.parse_args()

    main(args.template, args.data, args.out, args.header)
